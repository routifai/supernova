"""A spreadsheet or CSV file as a schema, never as rows of text.

Numbers in a table are answered by code (pandas, duckdb) in the Computer, not by reading rows as
text, so the index holds no row of a CSV or workbook. What it keeps instead:

* a **profile** per sheet: where the header row really is, the columns with an inferred type, the
  row count, the raw rows above the header and the first few rows below it. It is the manifest a
  model reads (``files_get``, an attachment's context), written as Markdown with one page per
  sheet;
* **schema cards**: the file name, its sheet names and *every* column name, cut into passages, so
  the search can find "which file has a ``revenue`` column?".

The file is read as a stream and only so far: a cap on the rows scanned, a cap on the unpacked size
of a workbook. Pure apart from reading the file; nothing here calls a model.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from pathlib import Path

from omnigent.runner.knowledge.chunking import MAX_CHARS, Block, Chunk, ParsedPage
from omnigent.runner.knowledge.fsio import open_nofollow
from omnigent.runner.knowledge.markdown import to_markdown

#: File kinds handled as tables.
TABULAR_KINDS = frozenset({"csv", "xlsx"})
#: The chunk kind a schema card is stored with.
SCHEMA_CHUNK_KIND = "schema"
PREVIEW_ROWS = 5
#: Rows read to infer a column's type.
TYPE_SAMPLE_ROWS = 1000
#: Rows counted before the count is reported as "at least".
MAX_SCAN_ROWS = 200_000
#: Rows from the top of a sheet searched for the header (a title or blank lines may precede it).
HEADER_SEARCH_ROWS = 20
MAX_SHEETS = 10
#: Columns shown with a type in the manifest; the rest are counted ("and N more").
MAX_SHOWN_COLUMNS = 60
#: Column names kept (all are searchable) and columns whose type is inferred.
MAX_COLUMN_NAMES = 2000
MAX_TYPED_COLUMNS = 200
PREVIEW_COLUMNS = 20
CELL_CHARS = 30
TOP_ROWS_SHOWN = 5
#: A workbook that unpacks to more than this is not opened (a zip bomb, or simply too big).
MAX_UNPACKED_BYTES = 512 * 1024 * 1024
MAX_ZIP_ENTRIES = 5000
#: Rows read in formula mode (the header, and enough rows to see a formula column).
MAX_FORMULA_ROWS = HEADER_SEARCH_ROWS + TYPE_SAMPLE_ROWS
FORMULA_TYPE = "formula (not computed)"

#: Longest field a CSV may hold (a cell can be a pasted document; the default 128 KB would fail
#: the whole file). Raised only while a file is being read, then put back.
CSV_FIELD_LIMIT = 16 * 1024 * 1024
#: Longest sheet or column name kept (a header cell can hold a paragraph).
MAX_NAME_CHARS = 120
#: Largest single unpacked part of a workbook that openpyxl loads whole (shared strings); the
#: sheets themselves are streamed.
MAX_SHARED_STRINGS_BYTES = 64 * 1024 * 1024

_INTEGER = re.compile(r"^[+-]?\d+$")
_THOUSANDS = re.compile(r"^[+-]?\d{1,3}(,\d{3})+(\.\d+)?$")
_NUMBER = re.compile(r"^[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?$")
_DECIMAL_COMMA = re.compile(r"^[+-]?(\d{1,3}(\.\d{3})+|\d+),\d+$")
_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}([ T]\d{2}:\d{2}(:\d{2})?)?$")
_BOOLEAN = frozenset({"true", "false"})
_BOMS = (
    (b"\xef\xbb\xbf", "utf-8-sig"),
    (b"\xff\xfe\x00\x00", "utf-32"),
    (b"\x00\x00\xfe\xff", "utf-32"),
    (b"\xff\xfe", "utf-16"),
    (b"\xfe\xff", "utf-16"),
)


@dataclass(frozen=True)
class Column:
    """One column: its header text and the type its values share."""

    name: str
    type: str


@dataclass(frozen=True)
class SheetProfile:
    """One sheet (a CSV file is one sheet).

    :param rows: Data rows below the header; see *rows_note* for how far to trust it.
    :param rows_note: ``exact``, ``at_least`` (the scan cap was reached) or ``dimension`` (the
        workbook's own idea of the sheet size, when the scan cap was reached).
    :param header_row: Position of the header among the sheet's non-empty rows (1 is the top).
    """

    name: str
    columns: list[Column]
    columns_total: int
    rows: int
    column_names: list[str] = field(default_factory=list)
    rows_note: str = "exact"
    header_row: int = 1
    top_rows: list[list[str]] = field(default_factory=list)
    preview: list[list[str]] = field(default_factory=list)


@dataclass(frozen=True)
class TableProfile:
    """What a table file holds, without its rows."""

    sheets: list[SheetProfile]
    sheet_names: list[str] = field(default_factory=list)

    @property
    def sheets_total(self) -> int:
        return max(len(self.sheet_names), len(self.sheets))


# ---------------------------------------------------------------------------- cells and types


def _name(value: object, fallback: str) -> str:
    """A sheet or column name: its text, cut to :data:`MAX_NAME_CHARS`, or *fallback*."""
    text = _text(value).replace("\n", " ")
    return (text if len(text) <= MAX_NAME_CHARS else text[: MAX_NAME_CHARS - 1] + "…") or fallback


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    return str(value).strip()


def _shown(value: object) -> str:
    text = _text(value).replace("\n", " ")
    return text if len(text) <= CELL_CHARS else text[: CELL_CHARS - 1] + "…"


def _shown_row(row: list[object]) -> list[str]:
    cells = [_shown(c) for c in row]
    while cells and not cells[-1]:
        cells.pop()
    return cells


def _type_of(value: object) -> str | None:
    """The type of one cell value (``None`` for an empty cell)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "integer" if value.is_integer() else "number"
    if isinstance(value, datetime | date):
        return "date"
    text = str(value).strip()
    if not text:
        return None
    if _INTEGER.match(text):
        # An id with leading zeros ("007") is text: reading it as a number would drop them.
        return (
            "text"
            if len(text.lstrip("+-")) > 1 and text.lstrip("+-").startswith("0")
            else "integer"
        )
    if _NUMBER.match(text) or _THOUSANDS.match(text):
        return "number"
    if _DECIMAL_COMMA.match(text):
        return "number (decimal comma)"
    if _DATE.match(text):
        return "date"
    if text.lower() in _BOOLEAN:
        return "boolean"
    return "text"


def _merge(seen: set[str]) -> str:
    """One column type from the types of its sampled values."""
    if not seen:
        return "empty"
    if len(seen) == 1:
        return next(iter(seen))
    if seen <= {"integer", "number"}:
        return "number"
    if seen <= {"integer", "number (decimal comma)"}:
        return "number (decimal comma)"
    return "text"


# ---------------------------------------------------------------------------- one sheet


def _filled(row: list[object]) -> int:
    return sum(1 for c in row if _text(c))


def _profile_rows(
    name: str,
    rows: Iterable[list[object]],
    *,
    dimension: int | None = None,
) -> tuple[SheetProfile, int]:
    """A sheet from its non-empty rows; also the position of its header among them (0-based).

    The header is the first of the top :data:`HEADER_SEARCH_ROWS` rows that has at least half as
    many filled cells as the widest of them (a title row above it has one or two).
    """
    iterator: Iterator[list[object]] = iter(rows)
    head: list[list[object]] = []
    for row in iterator:
        head.append(row)
        if len(head) >= HEADER_SEARCH_ROWS:
            break
    if not head:
        return SheetProfile(name, [], 0, 0), 0
    widest = max(_filled(r) for r in head)
    at = next(i for i, r in enumerate(head) if _filled(r) * 2 >= widest)
    header = head[at]
    count = max(len(header) - 1, 0)
    while count and not _text(header[count]):
        count -= 1  # trailing empty header cells are not columns
    width = count + 1 if header else 0
    names = [_name(header[i], f"column {i + 1}") for i in range(min(width, MAX_COLUMN_NAMES))]
    typed = min(len(names), MAX_TYPED_COLUMNS)
    seen: list[set[str]] = [set() for _ in range(typed)]
    preview: list[list[str]] = []
    total = 0
    capped = False

    def data() -> Iterator[list[object]]:
        yield from head[at + 1 :]
        yield from iterator

    for row in data():
        if total >= MAX_SCAN_ROWS:
            capped = True
            break
        total += 1
        if total <= TYPE_SAMPLE_ROWS:
            for index in range(typed):
                kind = _type_of(row[index]) if index < len(row) else None
                if kind:
                    seen[index].add(kind)
        if len(preview) < PREVIEW_ROWS:
            preview.append([_shown(row[i]) if i < len(row) else "" for i in range(len(names))])
    types = [_merge(s) for s in seen] + ["unknown"] * (len(names) - typed)
    note = "exact"
    if capped:
        note = "at_least"
        if dimension is not None and dimension > total + at + 1:
            total, note = dimension - at - 1, "dimension"
    top = [_shown_row(r) for r in head[:at]][:TOP_ROWS_SHOWN]
    profile = SheetProfile(
        name=name,
        columns=[Column(n, t) for n, t in zip(names, types, strict=True)][:MAX_SHOWN_COLUMNS],
        columns_total=len(names) if width <= MAX_COLUMN_NAMES else width,
        rows=total,
        column_names=names,
        rows_note=note,
        header_row=at + 1,
        top_rows=top,
        preview=preview,
    )
    return profile, at


# ---------------------------------------------------------------------------- CSV


def _encoding_of(sample: bytes) -> str:
    for bom, encoding in _BOMS:
        if sample.startswith(bom):
            return encoding
    try:
        sample.decode("utf-8")
    except UnicodeDecodeError as exc:
        if exc.start < len(sample) - 4:  # not just a character cut by the sample's end
            return "cp1252"
    return "utf-8"


def _csv_rows(path: Path) -> Iterator[list[object]]:
    with open_nofollow(path) as raw:
        encoding = _encoding_of(raw.read(65536))
        raw.seek(0)
        text = io.TextIOWrapper(raw, encoding=encoding, errors="replace", newline="")
        sample = text.read(4096)
        text.seek(0)
        try:
            dialect: type[csv.Dialect] = csv.Sniffer().sniff(  # type: ignore[assignment]
                sample.lstrip("﻿"), delimiters=",;\t|"
            )
        except csv.Error:
            dialect = csv.excel
        previous = csv.field_size_limit(CSV_FIELD_LIMIT)
        try:
            for row in csv.reader(text, dialect):
                if any(c.strip() for c in row):
                    yield [c.lstrip("\ufeff") if i == 0 else c for i, c in enumerate(row)]
        finally:
            csv.field_size_limit(previous)


# ---------------------------------------------------------------------------- XLSX


def _check_unpacked_size(path: Path) -> None:
    """Refuse a workbook that would unpack to an absurd size, before openpyxl inflates it."""
    try:
        with open_nofollow(path) as handle, zipfile.ZipFile(handle) as archive:
            infos = archive.infolist()
    except (zipfile.BadZipFile, OSError):
        raise ValueError("not a readable workbook") from None
    if len(infos) > MAX_ZIP_ENTRIES or sum(i.file_size for i in infos) > MAX_UNPACKED_BYTES:
        raise ValueError("This workbook is too large to read")
    if any(
        i.filename.endswith("sharedStrings.xml") and i.file_size > MAX_SHARED_STRINGS_BYTES
        for i in infos
    ):
        raise ValueError("This workbook has too much text to read")


def _formula_pass(path: Path, sheet: str, profile: SheetProfile, header_at: int) -> SheetProfile:
    """Re-read the header and the first rows of *sheet* with formulas.

    A file written by a library and never opened in Excel has formulas without computed values:
    the values pass sees empty cells, so a column of them (or a header made of one) would be lost
    or called empty. Such a column is named by its formula and typed ``formula (not computed)``.
    """
    from openpyxl import load_workbook

    with open_nofollow(path) as handle:
        book = load_workbook(handle, read_only=True, data_only=False)
        try:
            rows = [
                list(r) for r in book[sheet].iter_rows(max_row=MAX_FORMULA_ROWS, values_only=True)
            ]
        finally:
            book.close()
    rows = [r for r in rows if any(_text(c) for c in r)]
    if len(rows) <= header_at:
        return profile
    body = rows[header_at + 1 :]
    names = list(profile.column_names)
    types = [c.type for c in profile.columns]
    types += ["unknown"] * (len(names) - len(types))
    width = max(
        (max((i + 1 for i, c in enumerate(r) if _text(c)), default=0) for r in rows[header_at:]),
        default=0,
    )
    width = min(width, MAX_TYPED_COLUMNS)
    changed = False
    for i in range(width):
        cell = rows[header_at][i] if i < len(rows[header_at]) else None
        if i >= len(names):
            names.append(_name(cell, f"column {i + 1}"))
            types.append("empty")
            changed = True
        elif names[i].startswith("column ") and isinstance(cell, str) and cell.strip():
            names[i] = _name(cell, names[i])
            changed = True
        formula = any(i < len(r) and isinstance(r[i], str) and r[i].startswith("=") for r in body)
        if formula and types[i] == "empty":
            types[i] = FORMULA_TYPE
            changed = True
    if not changed:
        return profile
    columns = [Column(n, t) for n, t in zip(names, types, strict=True)]
    return replace(
        profile,
        columns=columns[:MAX_SHOWN_COLUMNS],
        column_names=names,
        columns_total=max(profile.columns_total, len(names)),
    )


def _xlsx_profile(path: Path) -> TableProfile:
    from openpyxl import load_workbook

    _check_unpacked_size(path)
    handle = open_nofollow(path)
    try:
        book = load_workbook(handle, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types for a bad file
        handle.close()
        raise ValueError(f"not a readable workbook ({type(exc).__name__})") from None
    sheets: list[SheetProfile] = []
    try:
        names = [_name(n, "Sheet") for n in book.sheetnames]
        positions: list[int] = []
        for sheet in book.worksheets[:MAX_SHEETS]:
            rows = (list(r) for r in sheet.iter_rows(values_only=True) if any(_text(c) for c in r))
            dimension = sheet.max_row if isinstance(sheet.max_row, int) else None
            profile, at = _profile_rows(_name(sheet.title, "Sheet"), rows, dimension=dimension)
            sheets.append(profile)
            positions.append(at)
    finally:
        book.close()
        handle.close()
    sheets = [
        _formula_pass(path, p.name, p, at) if p.columns else p
        for p, at in zip(sheets, positions, strict=True)
    ]
    return TableProfile(sheets or [SheetProfile("Sheet1", [], 0, 0)], names)


def profile_table(kind: str, path: Path, *, name: str = "") -> TableProfile:
    """The profile of a CSV or XLSX file, read from *path* as a stream.

    :raises ValueError: the file cannot be read as that kind.
    """
    if kind == "xlsx":
        return _xlsx_profile(Path(path))
    if kind == "csv":
        profile, _ = _profile_rows(name or "data", _csv_rows(Path(path)))
        return TableProfile([profile], [profile.name])
    raise ValueError(f"{kind} is not a table")


# ---------------------------------------------------------------------------- what is kept


def _rows_text(sheet: SheetProfile) -> str:
    if sheet.rows_note == "at_least":
        return f"at least {sheet.rows}"
    if sheet.rows_note == "dimension":
        return f"about {sheet.rows} (the sheet's own size; not all were read)"
    return str(sheet.rows)


def manifest_pages(profile: TableProfile, file_name: str) -> list[ParsedPage]:
    """The manifest as pages (one per sheet) so it is written with the usual page markers."""
    pages: list[ParsedPage] = []
    total = profile.sheets_total
    for number, sheet in enumerate(profile.sheets, start=1):
        blocks = [Block(sheet.name, heading_level=1)]
        if number == 1:
            more = (
                f" (and {total - len(profile.sheets)} more not shown)"
                if total > len(profile.sheets)
                else ""
            )
            blocks.append(
                Block(
                    f"Table file {file_name}: {total} sheet(s){more}. The rows are not indexed "
                    "or shown as text; compute on the file with code (pandas, duckdb) in the "
                    "Computer."
                )
            )
        if not sheet.columns:
            blocks.append(Block("This sheet is empty."))
        else:
            cols = ", ".join(f"{c.name} ({c.type})" for c in sheet.columns)
            more_cols = sheet.columns_total - len(sheet.columns)
            suffix = f", and {more_cols} more columns" if more_cols > 0 else ""
            lines = [f"Rows: {_rows_text(sheet)}. Columns ({sheet.columns_total}): {cols}{suffix}"]
            if sheet.header_row > 1:
                lines.append(f"The header is row {sheet.header_row}; the rows above it:")
                lines += [" | ".join(r) for r in sheet.top_rows]
            blocks.append(Block("\n".join(lines)))
            shown = sheet.columns[:PREVIEW_COLUMNS]
            header = " | ".join(c.name for c in shown)
            rows = [header, *(" | ".join(row[: len(shown)]) for row in sheet.preview)]
            blocks.append(Block("\n".join(rows), atomic=True))
        pages.append(ParsedPage(number, blocks))
    return pages


def manifest_markdown(profile: TableProfile, file_name: str) -> str:
    """The manifest a model reads in place of the file's rows."""
    return to_markdown(manifest_pages(profile, file_name), kind="sheet")


def schema_cards(profile: TableProfile, file_name: str) -> list[Chunk]:
    """The passages the index keeps: file name, every sheet name and every column name."""
    lines = [
        f"Table file: {file_name}",
        "Sheets: " + ", ".join(profile.sheet_names or [s.name for s in profile.sheets]),
    ]
    for sheet in profile.sheets:
        names = sheet.column_names
        extra = sheet.columns_total - len(names)
        tail = f", and {extra} more" if extra > 0 else ""
        lines.append(f"Sheet {sheet.name} columns: " + ", ".join(names) + tail)
    return [Chunk(1, 1, SCHEMA_CHUNK_KIND, "", piece) for piece in _cut(lines)]


def _cut(lines: list[str]) -> Iterator[str]:
    """The lines as pieces of at most :data:`MAX_CHARS`, cut at ", " when a line is longer."""
    current = ""
    for line in lines:
        while len(line) > MAX_CHARS:
            cut = line.rfind(", ", 0, MAX_CHARS)
            cut = cut + 2 if cut > MAX_CHARS // 2 else MAX_CHARS
            head, line = line[:cut], line[cut:]
            if current:
                yield current
                current = ""
            yield head.rstrip()
        if current and len(current) + 1 + len(line) > MAX_CHARS:
            yield current
            current = ""
        current = f"{current}\n{line}" if current else line
    if current:
        yield current
