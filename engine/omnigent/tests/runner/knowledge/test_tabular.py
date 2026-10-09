"""Tables are a schema, never rows of text: the manifest, the card, and what search finds."""

from __future__ import annotations

import io
import os
import threading
import time
from pathlib import Path

import pytest

from omnigent.runner.knowledge.db import STATE_QUEUED, STATE_READY
from omnigent.runner.knowledge.indexer import Indexer
from omnigent.runner.knowledge.markdown import split_pages, to_markdown
from omnigent.runner.knowledge.tabular import profile_table
from tests.runner.knowledge._fixtures import FakeEmbedder, build_indexer, run_all

UPLOADS = "your_files/uploads/2026-10-09"
CSV = "id,region,revenue,day\n" + "\n".join(
    f"{i},north,{i * 1.5},2026-01-{i % 28 + 1:02d}" for i in range(1, 41)
)


def _put(workspace: Path, rel: str, data: bytes | str) -> None:
    path = workspace / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data if isinstance(data, bytes) else data.encode())


def _xlsx() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "Sales"
    book.active.append(["region", "total", "shipped"])
    book.active.append(["north", 12.5, True])
    book.active.append(["south", 7, False])
    book.create_sheet("Costs").append(["item", "cost"])
    book["Costs"].append(["rent", 5])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _file(tmp_path: Path, name: str, data: bytes | str) -> Path:
    path = tmp_path / name
    path.write_bytes(data if isinstance(data, bytes) else data.encode())
    return path


def test_csv_profile_has_columns_types_rows_and_a_five_row_preview(tmp_path: Path) -> None:
    (sheet,) = profile_table("csv", _file(tmp_path, "s.csv", CSV), name="s.csv").sheets
    assert [(c.name, c.type) for c in sheet.columns] == [
        ("id", "integer"),
        ("region", "text"),
        ("revenue", "number"),
        ("day", "date"),
    ]
    assert sheet.rows == 40 and len(sheet.preview) == 5 and sheet.preview[0][:2] == ["1", "north"]


def test_xlsx_profile_is_per_sheet(tmp_path: Path) -> None:
    sales, costs = profile_table("xlsx", _file(tmp_path, "b.xlsx", _xlsx())).sheets
    assert (sales.name, sales.rows) == ("Sales", 2)
    assert [(c.name, c.type) for c in sales.columns] == [
        ("region", "text"),
        ("total", "number"),
        ("shipped", "boolean"),
    ]
    assert (costs.name, costs.rows, costs.columns[1].type) == ("Costs", 1, "integer")


def test_ingest_of_a_csv_is_the_manifest_and_no_row_is_indexed(tmp_path: Path) -> None:
    indexer, ws = build_indexer(tmp_path)
    _put(ws, f"{UPLOADS}/sales.csv", CSV)
    got = indexer.ingest(f"{UPLOADS}/sales.csv", inline=True)
    assert got["kind"] == "csv" and got["pages"] == 1
    text = got["markdown"]
    assert "id (integer), region (text), revenue (number), day (date)" in text
    assert "Rows: 40" in text and "pandas" in text
    assert "| 1 | north | 1.5 | 2026-01-02 |" in text and "| 6 |" not in text  # 5 rows only
    chunks = indexer.db.chunks_for(got["file_id"])
    assert len(chunks) == 1 and chunks[0].kind == "schema"  # no row-level chunks
    assert "north" not in chunks[0].text and "revenue" in chunks[0].text


def test_search_finds_a_table_by_a_column_name_but_not_by_a_value(tmp_path: Path) -> None:
    indexer, ws = build_indexer(tmp_path)
    _put(ws, f"{UPLOADS}/sales.csv", CSV)
    _put(ws, f"{UPLOADS}/book.xlsx", _xlsx())
    indexer.scan()
    run_all(indexer)
    hit = indexer.search("revenue", mode="keyword")["results"][0]
    assert hit["file_name"] == "sales.csv" and "revenue" in hit["passage"]
    assert indexer.search("shipped", mode="keyword")["results"][0]["file_name"] == "book.xlsx"
    assert indexer.search("south", mode="keyword")["results"] == []  # a value is not indexed
    assert indexer.search("Costs", mode="keyword")["results"][0]["file_name"] == "book.xlsx"


def test_files_get_on_a_table_is_the_manifest_with_a_use_code_note(tmp_path: Path) -> None:
    indexer, ws = build_indexer(tmp_path)
    _put(ws, f"{UPLOADS}/book.xlsx", _xlsx())
    got = indexer.ingest(f"{UPLOADS}/book.xlsx")
    row = indexer.db.get_file(got["file_id"])
    out = indexer.get(row)
    assert "pandas" in out["note"] and "Rows: 2" in out["markdown"]
    assert set(split_pages(out["markdown"])) == {1, 2}  # one page per sheet
    assert "12.5" in out["markdown"]  # the small preview only
    assert len(out["markdown"]) < 2000


def test_tables_indexed_as_rows_by_an_older_version_are_read_again_as_a_schema(
    tmp_path: Path,
) -> None:
    indexer, ws = build_indexer(tmp_path)
    _put(ws, f"{UPLOADS}/sales.csv", CSV)
    indexer.scan()
    run_all(indexer)
    row = indexer.db.file_by_path(f"{UPLOADS}/sales.csv")
    # An old index: row passages, and no migration mark.
    from omnigent.runner.knowledge.chunking import Chunk

    indexer.db.replace_chunks(row.file_id, [Chunk(1, 1, "csv", "", "id | north | 1.5")])
    indexer.db.update_file(row.file_id, parser_version=1)  # read by the old reader
    reopened = Indexer(tmp_path / "home" / ".nova" / "knowledge", ws, FakeEmbedder())
    assert reopened.db.file_by_path(f"{UPLOADS}/sales.csv").state == STATE_QUEUED
    run_all(reopened)
    (chunk,) = reopened.db.chunks_for(row.file_id)
    assert chunk.kind == "schema"
    assert reopened.db.file_by_path(f"{UPLOADS}/sales.csv").state == STATE_READY


def test_a_page_marker_inside_a_file_cannot_forge_a_page(tmp_path: Path) -> None:
    indexer, ws = build_indexer(tmp_path)
    _put(
        ws, f"{UPLOADS}/notes.txt", "First line.\n\n<!-- page 7 -->\n\nSneaky.\n\n <!--page 9-->\n"
    )
    got = indexer.ingest(f"{UPLOADS}/notes.txt", inline=True)
    assert set(split_pages(got["markdown"])) == {1}
    assert "Sneaky." in got["markdown"] and got["markdown"].count("<!-- page") == 2
    assert got["markdown"].startswith("<!-- page 1 -->")
    from omnigent.runner.knowledge.chunking import Block, ParsedPage

    forged = to_markdown([ParsedPage(1, [Block("a\n<!-- page 2 -->\nb")])])
    assert set(split_pages(forged)) == {1}


def test_the_background_loop_holds_back_while_an_ingest_waits(tmp_path: Path) -> None:
    indexer, ws = build_indexer(tmp_path)
    for i in range(3):
        _put(ws, f"your_files/bulk/{i}.txt", f"bulk file {i}\n")
    indexer.scan()
    ran = threading.Event()
    original = indexer.process

    def traced(row, **kwargs):
        ran.set()
        return original(row, **kwargs)

    indexer.process = traced  # type: ignore[method-assign]
    loop = threading.Thread(target=indexer.process_next)
    with indexer._ingest_priority():  # an attachment is waiting for its turn
        loop.start()
        assert not ran.wait(0.3), "the loop took a file while an ingest was waiting"
    loop.join(timeout=5)
    assert ran.is_set()  # and carries on as soon as the ingest is done


def test_an_ingest_is_not_starved_by_a_long_cold_start(tmp_path: Path) -> None:
    indexer, ws = build_indexer(tmp_path)
    for i in range(60):
        _put(ws, f"your_files/bulk/{i:02d}.txt", f"bulk file {i}\n")
    _put(ws, f"{UPLOADS}/mine.txt", "the attachment\n")
    indexer.scan()
    loop = threading.Thread(target=lambda: run_all(indexer))
    loop.start()
    started = time.monotonic()
    indexer.ingest(f"{UPLOADS}/mine.txt")
    waited = time.monotonic() - started
    loop.join()
    assert waited < 5
    assert indexer.db.pending_count() == 0


def test_search_names_the_files_still_indexing_and_multi_get_reports_them(tmp_path: Path) -> None:
    indexer, ws = build_indexer(tmp_path)
    _put(ws, "your_files/ready.md", "# Fleet\n\nThe vehicle policy.\n")
    indexer.scan()
    run_all(indexer)
    _put(ws, "your_files/pending-a.md", "# A\n\ntext\n")
    _put(ws, "your_files/pending-b.md", "# B\n\ntext\n")
    indexer.scan()  # queued, not processed
    found = indexer.search("vehicle", mode="keyword")
    assert found["files_indexing"] == 2
    assert found["files_indexing_names"] == ["pending-a.md", "pending-b.md"]
    got = indexer.multi_get("*.md")
    assert got["returned"] == 1
    assert [f["path"] for f in got["not_ready"]] == [
        "your_files/pending-a.md",
        "your_files/pending-b.md",
    ]
    assert "still being indexed" in got["note"]
    run_all(indexer)
    assert "not_ready" not in indexer.multi_get("*.md")


# ---------------------------------------------------------------------------- real sheets


def _book(tmp_path: Path, build) -> Path:
    from openpyxl import Workbook

    book = Workbook()
    build(book)
    path = tmp_path / "book.xlsx"
    book.save(path)
    return path


def test_the_header_is_found_below_a_title_and_the_raw_top_rows_are_kept(tmp_path: Path) -> None:
    def build(book) -> None:
        sheet = book.active
        sheet.append(["Quarterly report"])
        sheet.append([])
        sheet.append(["Confidential"])
        sheet.append(["region", "units", "price", "day"])
        sheet.append(["north", 3, 2.5, "2026-01-01"])

    profile = profile_table("xlsx", _book(tmp_path, build))
    (sheet,) = profile.sheets
    assert [c.name for c in sheet.columns] == ["region", "units", "price", "day"]
    assert sheet.header_row == 3 and sheet.rows == 1
    assert sheet.top_rows == [["Quarterly report"], ["Confidential"]]
    from omnigent.runner.knowledge.tabular import manifest_markdown

    text = manifest_markdown(profile, "book.xlsx")
    assert "Quarterly report" in text and "The header is row 3" in text


def test_a_wide_sheet_reports_its_real_size_and_every_name_stays_searchable(
    tmp_path: Path,
) -> None:
    names = [f"metric_{i:03d}" for i in range(300)]

    def build(book) -> None:
        book.active.append(names)
        book.active.append(list(range(300)))
        for i in range(14):
            book.create_sheet(f"extra {i}").append(["a"])

    profile = profile_table("xlsx", _book(tmp_path, build))
    assert len(profile.sheets) == 10 and profile.sheets_total == 15
    wide = profile.sheets[0]
    assert wide.columns_total == 300 and len(wide.columns) == 60
    from omnigent.runner.knowledge.tabular import manifest_markdown, schema_cards

    text = manifest_markdown(profile, "book.xlsx")
    assert "Columns (300)" in text and "and 240 more columns" in text
    assert "and 5 more not shown" in text and "15 sheet(s)" in text
    cards = "\n".join(c.text for c in schema_cards(profile, "book.xlsx"))
    assert "metric_299" in cards and "extra 13" in cards  # beyond what the manifest shows
    assert all(len(c.text) <= 1600 for c in schema_cards(profile, "book.xlsx"))


def test_formula_columns_without_cached_values_are_named_not_called_empty(tmp_path: Path) -> None:
    def build(book) -> None:
        sheet = book.active
        sheet.append(["item", "qty", "price", '=A1&"-total"'])
        for i in range(1, 4):
            sheet.append([f"x{i}", i, 2.0, f"=B{i + 1}*C{i + 1}"])

    (sheet,) = profile_table("xlsx", _book(tmp_path, build)).sheets
    types = {c.name: c.type for c in sheet.columns}
    assert types["item"] == "text" and types["qty"] == "integer"
    assert types['=A1&"-total"'] == "formula (not computed)"


def test_a_utf16_csv_is_read_with_its_bom(tmp_path: Path) -> None:
    raw = "id;région;montant\n1;nord;1,5\n2;sud;2,25\n".encode("utf-16")
    (sheet,) = profile_table("csv", _file(tmp_path, "u.csv", raw), name="u.csv").sheets
    assert [c.name for c in sheet.columns] == ["id", "région", "montant"]
    assert sheet.columns[2].type == "number (decimal comma)" and sheet.rows == 2


def test_leading_zero_ids_stay_text_and_big_cells_do_not_break_the_read(tmp_path: Path) -> None:
    big = "x" * 300_000
    data = f"zip,code,notes\n00123,A,{big}\n04567,B,short\n"
    (sheet,) = profile_table("csv", _file(tmp_path, "z.csv", data), name="z.csv").sheets
    assert [(c.name, c.type) for c in sheet.columns][:2] == [("zip", "text"), ("code", "text")]
    assert sheet.rows == 2 and sheet.preview[0][0] == "00123"


def test_a_csv_is_streamed_and_its_row_count_is_capped(tmp_path: Path, monkeypatch) -> None:
    from omnigent.runner.knowledge import tabular

    monkeypatch.setattr(tabular, "MAX_SCAN_ROWS", 500)
    path = _file(tmp_path, "big.csv", "a,b\n" + "\n".join(f"{i},x" for i in range(5000)))
    (sheet,) = profile_table("csv", path, name="big.csv").sheets
    assert (sheet.rows, sheet.rows_note) == (500, "at_least")
    assert "at least 500" in tabular.manifest_markdown(TableProfileOf(sheet), "big.csv")


def TableProfileOf(sheet):
    from omnigent.runner.knowledge.tabular import TableProfile

    return TableProfile([sheet], [sheet.name])


def test_a_workbook_that_unpacks_to_too_much_is_refused_before_it_is_opened(
    tmp_path: Path, monkeypatch
) -> None:
    from omnigent.runner.knowledge import tabular

    monkeypatch.setattr(tabular, "MAX_UNPACKED_BYTES", 1000)
    with pytest.raises(ValueError, match="too large"):
        profile_table("xlsx", _file(tmp_path, "b.xlsx", _xlsx()))


def test_a_value_is_not_a_zero_sum_trap_zeros_csv(tmp_path: Path) -> None:
    (sheet,) = profile_table(
        "csv", _file(tmp_path, "zeros.csv", "id,amount\n001,10\n002,20\n010,30\n"), name="z.csv"
    ).sheets
    assert sheet.columns[0].type == "text" and sheet.columns[1].type == "integer"


# ---------------------------------------------------------------------------- hashes and find


def test_find_upload_matches_by_content_and_rehashes_on_disk(tmp_path: Path) -> None:
    import hashlib

    indexer, ws = build_indexer(tmp_path)
    data = b"the same bytes\n"
    _put(ws, f"{UPLOADS}/a.txt", data)
    _put(ws, "your_files/elsewhere.txt", data)  # not an upload: never offered
    indexer.scan()
    run_all(indexer)
    digest = hashlib.sha256(data).hexdigest()
    found = indexer.find_upload(digest)
    assert found and found["path"] == f"{UPLOADS}/a.txt"
    assert indexer.find_upload(hashlib.sha256(b"other").hexdigest()) is None
    # The file changed after it was indexed (same size and mtime preserved): the stored hash is
    # only a lead, so the answer must be no.
    path = ws / UPLOADS / "a.txt"
    stat = path.stat()
    path.write_bytes(b"different byte\n")
    os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns))
    assert indexer.find_upload(digest) is None


def test_the_hash_is_kept_for_files_indexed_before_hashes_existed(tmp_path: Path) -> None:
    import hashlib

    indexer, ws = build_indexer(tmp_path)
    _put(ws, f"{UPLOADS}/a.txt", "old\n")
    indexer.scan()
    run_all(indexer)
    row = indexer.db.file_by_path(f"{UPLOADS}/a.txt")
    indexer.db.update_file(row.file_id, sha256=None)
    indexer.scan()
    assert (
        indexer.db.file_by_path(f"{UPLOADS}/a.txt").sha256 == hashlib.sha256(b"old\n").hexdigest()
    )


def test_only_files_read_by_an_older_reader_are_read_again(tmp_path: Path) -> None:
    indexer, ws = build_indexer(tmp_path)
    _put(ws, f"{UPLOADS}/a.txt", "text\n")
    _put(ws, f"{UPLOADS}/t.csv", CSV)
    indexer.scan()
    run_all(indexer)
    for name in ("a.txt", "t.csv"):
        assert indexer.db.file_by_path(f"{UPLOADS}/{name}").parser_version >= 1
    reopened = Indexer(tmp_path / "home" / ".nova" / "knowledge", ws, FakeEmbedder())
    assert {r.state for r in reopened.db.files()} == {STATE_READY}  # nothing outdated: no rework
    csv_row = reopened.db.file_by_path(f"{UPLOADS}/t.csv")
    reopened.db.update_file(csv_row.file_id, parser_version=1)
    again = Indexer(tmp_path / "home" / ".nova" / "knowledge", ws, FakeEmbedder())
    assert again.db.file_by_path(f"{UPLOADS}/t.csv").state == STATE_QUEUED
    assert again.db.file_by_path(f"{UPLOADS}/a.txt").state == STATE_READY


def test_a_swapped_in_symlink_is_not_followed(tmp_path: Path) -> None:
    from omnigent.runner.knowledge.fsio import open_nofollow

    target = tmp_path / "secret.txt"
    target.write_text("secret")
    link = tmp_path / "link.txt"
    link.symlink_to(target)
    with pytest.raises(OSError):
        open_nofollow(link)
    with open_nofollow(target) as handle:
        assert handle.read() == b"secret"


def test_stopping_the_service_ends_the_wait_for_an_ingest(tmp_path: Path) -> None:
    indexer, _ = build_indexer(tmp_path)
    done = threading.Event()
    with indexer._ingest_priority():
        waiter = threading.Thread(target=lambda: (indexer._yield_to_ingest(), done.set()))
        waiter.start()
        assert not done.wait(0.2)
        indexer.shutdown()
        assert done.wait(3)
    waiter.join()


def test_an_ingest_slips_in_between_the_text_pass_and_the_embedding_pass(tmp_path: Path) -> None:
    indexer, ws = build_indexer(tmp_path)
    _put(ws, "your_files/bulk.txt", "bulk words\n")
    indexer.scan()
    row = indexer.db.file_by_path("your_files/bulk.txt")
    background = threading.Thread(target=lambda: indexer.process(row, background=True))
    with indexer._ingest_priority():  # an attachment is waiting for its turn
        background.start()
        deadline = time.monotonic() + 3
        while time.monotonic() < deadline:
            if indexer.db.file_by_path("your_files/bulk.txt").state == "text_ready":
                break
            time.sleep(0.02)
        time.sleep(0.2)
        # the text is read, the embedding pass is held back for the waiting ingest
        assert indexer.db.file_by_path("your_files/bulk.txt").state == "text_ready"
        assert indexer._work.acquire(blocking=False)  # and the work lock is free for it
        indexer._work.release()
    background.join(5)
    assert indexer.db.file_by_path("your_files/bulk.txt").state == STATE_READY


def test_a_failed_file_is_not_queued_again_by_a_restart(tmp_path: Path) -> None:
    indexer, ws = build_indexer(tmp_path)
    _put(ws, f"{UPLOADS}/t.csv", CSV)
    indexer.scan()
    row = indexer.db.file_by_path(f"{UPLOADS}/t.csv")
    indexer.db.update_file(row.file_id, state="failed", parser_version=1, last_error="boom")
    again = Indexer(tmp_path / "home" / ".nova" / "knowledge", ws, FakeEmbedder())
    assert again.db.file_by_path(f"{UPLOADS}/t.csv").state == "failed"


def test_long_names_are_cut_and_the_shared_strings_part_is_capped(
    tmp_path: Path, monkeypatch
) -> None:
    from omnigent.runner.knowledge import tabular

    def build(book) -> None:
        book.active.append(["x" * 5000, "ok"])
        book.active.append(["a", "b"])

    path = _book(tmp_path, build)
    (sheet,) = profile_table("xlsx", path).sheets
    assert len(sheet.column_names[0]) == tabular.MAX_NAME_CHARS and sheet.column_names[1] == "ok"
    # a workbook whose shared-strings part is too big is refused before it is opened
    import zipfile

    big = tmp_path / "strings.xlsx"
    with zipfile.ZipFile(big, "w") as archive:
        archive.writestr("xl/sharedStrings.xml", "s" * 500)
    monkeypatch.setattr(tabular, "MAX_SHARED_STRINGS_BYTES", 100)
    with pytest.raises(ValueError, match="too much text"):
        profile_table("xlsx", big)


def test_the_csv_field_limit_is_raised_only_while_a_file_is_read(tmp_path: Path) -> None:
    import csv

    before = csv.field_size_limit()
    path = _file(tmp_path, "w.csv", "a,b\n" + "y" * 200_000 + ",1\n")
    (sheet,) = profile_table("csv", path, name="w.csv").sheets
    assert sheet.rows == 1
    assert csv.field_size_limit() == before
