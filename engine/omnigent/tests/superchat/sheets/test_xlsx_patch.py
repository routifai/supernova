"""Lossless XLSX cell edits: only target cells change, every other part survives byte-for-byte."""

from __future__ import annotations

import io
import re
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font

from omnigent.superchat.sheets import table as tbl
from omnigent.superchat.sheets.xlsx_patch import XlsxPatchError, patch_cells

_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d49444154789c6360f8cfc0f01f00050001ff89993d1d0000000049454e44ae426082"
)
_SPARK = (
    b'<extLst><ext uri="{05C60535-1F16-4fd2-B633-F4F36F0B64E0}" xmlns:x14="http://schemas.micro'
    b'soft.com/office/spreadsheetml/2009/9/main"><x14:sparklineGroups xmlns:xm="http://schemas.'
    b'microsoft.com/office/excel/2006/main"><x14:sparklineGroup><x14:sparklines><x14:sparkline>'
    b"<xm:f>Sales!B2:C2</xm:f><xm:sqref>D2</xm:sqref></x14:sparkline></x14:sparklines>"
    b"</x14:sparklineGroup></x14:sparklineGroups></ext></extLst>"
)
_PIVOT_CT = (
    b'<Override PartName="/xl/pivotTables/pivotTable1.xml" ContentType="application/vnd.openxml'
    b'formats-officedocument.spreadsheetml.pivotTable+xml"/>'
)


def _rich_workbook(*, with_calc_chain: bool = False) -> bytes:
    """Styles, a chart, an image, sparklines and a pivot part, as Excel would hold them."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Q3", "Q4"])
    for i in range(1, 6):
        ws.append([f"r{i}", 100 * i, f"=B{i + 1}*2"])
        ws.cell(row=i + 1, column=2).number_format = "#,##0.00"
        ws.cell(row=i + 1, column=2).font = Font(bold=True)
        ws.cell(row=i + 1, column=3).font = Font(italic=True)
    ws["B8"] = "=SUM(B2:B6)"
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=6), titles_from_data=True)
    ws.add_chart(chart, "E2")
    ws.add_image(XLImage(io.BytesIO(_PNG)), "E20")
    ws.freeze_panes = "A2"
    wb.create_sheet("Notes")["A1"] = "hello"
    buf = io.BytesIO()
    wb.save(buf)
    src = zipfile.ZipFile(buf)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "[Content_Types].xml":
                extra = _PIVOT_CT
                if with_calc_chain:
                    extra += (
                        b'<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.'
                        b'openxmlformats-officedocument.spreadsheetml.calcChain+xml"/>'
                    )
                data = data.replace(b"</Types>", extra + b"</Types>")
            elif info.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"</worksheet>", _SPARK + b"</worksheet>")
            elif info.filename == "xl/_rels/workbook.xml.rels" and with_calc_chain:
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/office'
                    b'Document/2006/relationships/calcChain" Target="calcChain.xml"/>'
                    b"</Relationships>",
                )
            z.writestr(info, data)
        z.writestr(
            "xl/pivotTables/pivotTable1.xml",
            b'<pivotTableDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/'
            b'2006/main" name="PT1" cacheId="1"/>',
        )
        if with_calc_chain:
            z.writestr(
                "xl/calcChain.xml",
                b'<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                b'<c r="C2" i="1"/><c r="C3"/></calcChain>',
            )
    return out.getvalue()


def _parts(raw: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        return {i.filename: z.read(i.filename) for i in z.infolist()}


def _cell(xml: bytes, ref: str) -> str:
    m = re.search(rf'<c r="{ref}"[^>]*?(?:/>|>.*?</c>)', xml.decode())
    assert m, f"{ref} missing"
    return m.group(0)


def test_only_sheet_and_workbook_parts_change_and_extras_survive() -> None:
    src = _rich_workbook()
    new, _ = tbl.apply_edits(
        src, "xlsx",
        [tbl.CellEdit("Sales", 1, 1, "1,234"), tbl.CellEdit("Sales", 2, 2, "=B3+1")],
    )  # fmt: skip
    before, after = _parts(src), _parts(new)
    assert list(before) == list(after)  # same entries, same order
    changed = {n for n in before if before[n] != after[n]}
    assert changed == {"xl/worksheets/sheet1.xml", "xl/workbook.xml"}
    names = set(after)
    assert any(n.startswith("xl/charts/chart") for n in names)
    assert any(n.startswith("xl/media/") for n in names)
    assert "xl/pivotTables/pivotTable1.xml" in names
    sheet = after["xl/worksheets/sheet1.xml"].decode()
    assert "<drawing " in sheet and "x14:sparklineGroups" in sheet  # references kept
    assert (
        after["xl/drawings/_rels/drawing1.xml.rels"]
        == before["xl/drawings/_rels/drawing1.xml.rels"]
    )


def test_edited_cells_new_values_and_styles_preserved() -> None:
    src = _rich_workbook()
    old_style = re.search(
        r'<c r="B2"[^>]*\bs="(\d+)"', _parts(src)["xl/worksheets/sheet1.xml"].decode()
    )
    new, summary = tbl.apply_edits(
        src, "xlsx",
        [tbl.CellEdit("Sales", 1, 1, "1,234.5"), tbl.CellEdit("Sales", 2, 2, "=B3+1"),
         tbl.CellEdit("Sales", 1, 0, "A & <b>"), tbl.CellEdit("Sales", 3, 0, True)],
    )  # fmt: skip
    xml = _parts(new)["xl/worksheets/sheet1.xml"]
    b2 = _cell(xml, "B2")
    assert f's="{old_style.group(1)}"' in b2 and "<v>1234.5</v>" in b2
    c3 = _cell(xml, "C3")
    assert "<f>B3+1</f>" in c3 and "<v>" not in c3 and 's="' in c3
    assert 't="inlineStr"' in _cell(xml, "A2") and "A &amp; &lt;b&gt;" in _cell(xml, "A2")
    assert 't="b"' in _cell(xml, "A4")
    ws = load_workbook(io.BytesIO(new))["Sales"]
    assert ws["B2"].value == 1234.5 and ws["C3"].value == "=B3+1" and ws["A2"].value == "A & <b>"
    assert ws["B2"].number_format == "#,##0.00" and ws["B2"].font.bold is True
    assert ws["B3"].value == 200 and ws["A1"].value == "Region"  # untouched cells
    assert "B2 100 → 1,234.5" in summary


def test_full_calc_on_load_set_and_updated() -> None:
    new, _ = tbl.apply_edits(_rich_workbook(), "xlsx", [tbl.CellEdit("Sales", 0, 0, "x")])
    wbxml = _parts(new)["xl/workbook.xml"].decode()
    assert len(re.findall(r'fullCalcOnLoad="1"', wbxml)) == 1
    again, _ = tbl.apply_edits(new, "xlsx", [tbl.CellEdit("Sales", 0, 0, "y")])
    assert _parts(again)["xl/workbook.xml"].decode().count("fullCalcOnLoad") == 1
    # An existing calcPr without the flag gets it; an explicit "0" is replaced.
    for calc in (b'<calcPr calcId="1"/>', b'<calcPr fullCalcOnLoad="0" calcId="1"/>'):
        parts = _parts(new)
        parts["xl/workbook.xml"] = re.sub(rb"<calcPr[^>]*/>", calc, parts["xl/workbook.xml"])
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            for n, d in parts.items():
                z.writestr(n, d)
        out = patch_cells(buf.getvalue(), [("Sales", 0, 0, "z")])
        text = _parts(out)["xl/workbook.xml"].decode()
        assert 'fullCalcOnLoad="1"' in text and 'fullCalcOnLoad="0"' not in text
        assert 'calcId="1"' in text


def test_clear_keeps_style_and_creates_rows_and_cells_in_order() -> None:
    src = _rich_workbook()
    new, _ = tbl.apply_edits(
        src, "xlsx",
        [tbl.CellEdit("Sales", 1, 1, ""), tbl.CellEdit("Sales", 9, 3, 7),
         tbl.CellEdit("Sales", 9, 1, 5), tbl.CellEdit("Sales", 6, 0, "gap"),
         tbl.CellEdit("Sales", 1, 5, 1)],
    )  # fmt: skip
    xml = _parts(new)["xl/worksheets/sheet1.xml"]
    b2 = _cell(xml, "B2")
    assert 's="' in b2 and "<v>" not in b2 and "<f>" not in b2
    rows = [int(r) for r in re.findall(r'<row r="(\d+)"', xml.decode())]
    assert rows == sorted(rows) and {7, 10} <= set(rows)
    row10 = re.search(r'<row r="10".*?</row>', xml.decode()).group(0)
    assert row10.index('r="B10"') < row10.index('r="D10"')
    row2 = re.search(r'<row r="2".*?</row>', xml.decode()).group(0)
    assert row2.index('r="B2"') < row2.index('r="F2"')
    assert 'ref="A1:F10"' in xml.decode()
    ws = load_workbook(io.BytesIO(new))["Sales"]
    assert ws["D10"].value == 7 and ws["B10"].value == 5 and ws["A7"].value == "gap"


def test_removing_a_formula_drops_calc_chain_cleanly() -> None:
    src = _rich_workbook(with_calc_chain=True)
    new, _ = tbl.apply_edits(src, "xlsx", [tbl.CellEdit("Sales", 1, 2, 5)])
    parts = _parts(new)
    assert "xl/calcChain.xml" not in parts
    assert b"calcChain" not in parts["[Content_Types].xml"]
    assert b"calcChain" not in parts["xl/_rels/workbook.xml.rels"]
    kept, _ = tbl.apply_edits(src, "xlsx", [tbl.CellEdit("Sales", 1, 1, 5)])
    assert "xl/calcChain.xml" in _parts(kept)  # no formula removed: the chain stays valid
    assert load_workbook(io.BytesIO(new))["Sales"]["C2"].value == 5


def test_unchanged_formula_keeps_cached_value_and_shared_formulas_expand() -> None:
    src = _rich_workbook()
    new, _ = tbl.apply_edits(src, "xlsx", [tbl.CellEdit("Sales", 1, 2, "=B2*2")])
    assert _parts(new)["xl/worksheets/sheet1.xml"] == _parts(src)["xl/worksheets/sheet1.xml"]
    # Shared formula: C2 master with C3 dependent; replacing the master keeps C3 working.
    parts = _parts(src)
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    sheet = re.sub(
        r'(<c r="C2"[^>]*>)<f>B2\*2</f>', r'\1<f t="shared" ref="C2:C3" si="0">B2*2</f>', sheet
    )
    sheet = re.sub(r'(<c r="C3"[^>]*>)<f>B3\*2</f>', r'\1<f t="shared" si="0"/>', sheet)
    parts["xl/worksheets/sheet1.xml"] = sheet.encode()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, d in parts.items():
            z.writestr(n, d)
    out = patch_cells(buf.getvalue(), [("Sales", 1, 2, 9)])
    ws = load_workbook(io.BytesIO(out))["Sales"]
    assert ws["C2"].value == 9 and ws["C3"].value == "=B3*2"


def test_errors() -> None:
    with pytest.raises(XlsxPatchError):
        patch_cells(_rich_workbook(), [("Nope", 0, 0, 1)])
    with pytest.raises(XlsxPatchError):
        patch_cells(b"not a zip", [("Sales", 0, 0, 1)])
    with pytest.raises(tbl.TableError):
        tbl.apply_edits(_rich_workbook(), "xlsx", [tbl.CellEdit("Nope", 0, 0, 1)])
